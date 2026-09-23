from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
p=Path(r'artifacts/agent-audit/audit-agent-recommandations.docx')
doc=Document(p)
# Remove the default Word title paragraph border that creates a blue rule.
style=doc.styles['Title']
ppr=style._element.get_or_add_pPr()
for bdr in list(ppr.findall(qn('w:pBdr'))): ppr.remove(bdr)
# Append implementation contract detail on two deliberately separated pages.
def page(): doc.add_page_break()
def h(s,l=1): doc.add_heading(s,l)
def para(s): doc.add_paragraph(s)
def bullet(s): doc.add_paragraph(s,style='List Bullet')
page(); h('Contrats techniques à préserver',1)
para('Cette section transforme les constats en contraintes de conception. Elle aide l’IA d’implémentation à corriger les écarts sans affaiblir les protections déjà présentes.')
h('Résolution du modèle planifié',2)
for s in [
'Calculer une seule fois l’identifiant effectif après application du fallback. La fiche effective doit alimenter le constructeur de modèle, le sélecteur d’outils, la fenêtre de contexte, les capacités de fichiers et les métadonnées persistées du run.',
'Ne pas utiliser le modèle demandé pour configurer le run après fallback. Les traces doivent conserver les deux identifiants (demandé et effectif) afin que l’opérateur comprenne pourquoi le modèle a changé.',
'Exercer les cas : modèle demandé absent du catalogue, modèle présent sans tools, modèle effectif avec contexte réduit et modèle par défaut indisponible. Chaque cas doit échouer explicitement ou produire un run cohérent.'
]: bullet(s)
h('Réservation interactive d’un run',2)
for s in [
'La lecture « run actif ? » suivie d’une insertion n’est pas un verrou. Choisir un mécanisme atomique compatible avec PostgreSQL et le statut multi-état ; un index unique partiel doit être étudié au regard des transitions waiting_for_user / waiting_for_approval.',
'Faire retourner au client un identifiant de run stable pour les répétitions de requête. Distinguer l’idempotence d’un retry réseau de la création d’une nouvelle tâche volontaire.',
'La requête perdante doit rejoindre le run existant ou recevoir une réponse de conflit explicite. Elle ne doit jamais lancer parallèlement une seconde boucle d’outils.'
]: bullet(s)
h('Contrat de validation des pièces jointes',2)
for s in [
'La validation doit être exécutée côté serveur après résolution des capacités effectives et avant le calcul du contexte, la création du run et tout appel fournisseur.',
'Le MIME fourni par le client est une indication, pas une preuve du contenu. Maintenir les garde-fous de lecture déjà présents : plafond d’octets, type/contenu cohérent, ZIP guard pour DOCX et fetch réseau défensif.',
'Les erreurs doivent nommer le type ou la capacité manquante, sans exposer l’URL interne ni les détails d’accès.'
]: bullet(s)
page(); h('Validation et critères de livraison',1)
h('Tests ciblés à ajouter',2)
t=doc.add_table(rows=1,cols=3); t.style='Table Grid'
for i,x in enumerate(['Zone','Scénario','Résultat attendu']):
 c=t.rows[0].cells[i]; c.text=x
 for r in c.paragraphs[0].runs:r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
 from docx.oxml import OxmlElement
 sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'17365D'); c._tc.get_or_add_tcPr().append(sh)
rows=[
('Scheduler/modèle','Modèle choisi sans outils, fallback avec outils','Sélection, contexte et run reposent tous sur le modèle effectif.'),
('Scheduler/modèle','Modèle de fallback absent ou invalide','Erreur exploitable avant création du run ; aucune tâche silencieusement dégradée.'),
('API/race','Deux requêtes POST même chat lancées simultanément','Un seul run actif et un seul ensemble d’exécutions d’outils.'),
('API/idempotence','Retry de la même requête après perte de réponse HTTP','Même run rendu, sans nouvel effet externe.'),
('Fichiers','PDF envoyé à un modèle sans capacité documents','Refus côté API avant run et avant appel du fournisseur.'),
('Fichiers','Image envoyée à un modèle sans vision','Refus explicite côté API ; aucune donnée d’image transmise.'),
('Fichiers','MIME incohérent avec le contenu téléchargé','Échec contrôlé et borné par les protections d’extraction existantes.'),
('Scheduler/timeout','Exécution dépasse la durée maximale du tick','Annulation propagée, checkpoint conservé, reprise idempotente.'),
('Scheduler/reprise','Run en attente d’approbation dépasse la lease','Politique explicite : attendre résolution ou expirer selon une durée métier ; pas de boucle cron infinie.')]
for row in rows:
 cells=t.add_row().cells
 for i,x in enumerate(row):
  cells[i].text=x
  for pp in cells[i].paragraphs: pp.paragraph_format.space_after=Pt(1)
  for rr in cells[i].paragraphs[0].runs: rr.font.size=Pt(8)
h('Séquence de livraison',2)
for s in [
'Lot A — correction locale du modèle effectif dans execute.ts et tests de contrat ; faible rayon d’impact, pas de migration.',
'Lot B — décision de réservation du run interactif : analyser les statuts actifs et la stratégie d’isolation, ajouter les tests de concurrence, puis seulement appliquer la migration si elle est nécessaire.',
'Lot C — validation MIME dans la route et tests de capacités ; préserver les modèles multimodaux valides.',
'Lot D — deadline de tick et politique de lease pour les états en attente ; instrumenter le résultat final et les reprises.',
'Lot E — observabilité et fonctionnalités Alpha, chacune derrière un flag si elle touche aux effets externes.'
]: bullet(s)
h('Définition de terminé',2)
para('Chaque correctif doit inclure un test déterministe du défaut, une vérification des scénarios existants d’approbation et de reprise, une migration vérifiée si le schéma change, et une trace d’observabilité qui permet de diagnostiquer le résultat en production sans journaliser de secret ou de contenu privé. La revue doit confirmer qu’aucun chemin interactif ou planifié ne continue à utiliser une fiche de modèle obsolète.')
# Set explicit black title and heading styles, without inherited borders.
for name in ['Title','Heading 1','Heading 2']:
 st=doc.styles[name]; st.font.color.rgb=RGBColor(0,0,0)
doc.save(p)
print(p)
